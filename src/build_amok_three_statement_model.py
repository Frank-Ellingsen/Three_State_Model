from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


ROOT = Path(__file__).resolve().parent.parent
WIDE_FILE = ROOT / "data" / "processed" / "amok_equipment_alle_aar_wide.csv"
CASH_FLOW_FILE = ROOT / "data" / "raw" / "amok_cash_flow_2024_2025.csv"
OUTPUT_FILE = ROOT / "models" / "Amok_Equipment_Integrated_3_Statement_Model.xlsx"
OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)


YEARS = list(range(2023, 2031))
ACTUAL_YEARS = YEARS[:3]
FORECAST_YEARS = YEARS[3:]
COLS = {year: year - 2021 for year in YEARS}  # B:I
SCENARIOS = ["Base Case", "Bull Case", "Bear Case", "Severe Stress"]

NAVY = "17365D"
BLUE = "4472C4"
DARK_BLUE = "2F5496"
LIGHT_BLUE = "DCE6F1"
LIGHT_GREEN = "E2EFDA"
LIGHT_RED = "FCE4D6"
PALE_YELLOW = "FFF2CC"
WHITE = "FFFFFF"
BLACK = "000000"
GRAY = "666666"
INPUT_BLUE = "0000FF"
FORMULA_GREEN = "008000"

FMT_NOK = '#,##0;[Red](#,##0);-'
FMT_NOK_1 = '#,##0.0;[Red](#,##0.0);-'
FMT_PCT = '0.0%;[Red](0.0%);-'
FMT_MULT = '0.0x'
FMT_DAYS = '0.0'
FMT_YEAR = '0'


def read_historical_records() -> tuple[dict[int, dict[str, float]], dict[int, dict[str, float]]]:
    financials: dict[int, dict[str, float]] = {}
    numeric_fields = {
        "salgsinntekter", "driftskostnader", "lonnskostnad", "driftsresultat",
        "netto_finans", "skattekostnad", "arsresultat", "anleggsmidler",
        "omlopsmidler", "varer", "fordringer", "bank", "sum_eiendeler",
        "kortsiktig_gjeld", "langsiktig_gjeld", "sum_gjeld", "egenkapital",
    }
    with WIDE_FILE.open("r", encoding="cp1252", newline="") as source:
        reader = csv.DictReader(source, delimiter=";")
        for record in reader:
            year = int(record["regnskapsaar"].lstrip("\ufeffï»¿"))
            financials[year] = {}
            for key, value in record.items():
                normalized_key = key.lstrip("\ufeffï»¿") if key else ""
                if not value or normalized_key not in numeric_fields:
                    continue
                financials[year][normalized_key] = float(value.replace(",", "."))

    cash_flows: dict[int, dict[str, float]] = {year: {} for year in ACTUAL_YEARS}
    with CASH_FLOW_FILE.open("r", encoding="cp1252", newline="") as source:
        reader = csv.DictReader(source)
        for record in reader:
            year = int(record["Kontantstrøm År"])
            label = record["Navn"]
            cash_flows[year][label] = float(record["Beløp"])
    return financials, cash_flows


def style_sheet(ws, tab_color: str = BLUE) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = "B5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.selection[0].activeCell = "B5"
    ws.sheet_view.selection[0].sqref = "B5"


def title(ws, text: str, last_col: int = 9) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=15)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 27
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A2"] = '=IF(Model_Checks!$B$2=0,"ALL CHECKS PASSED","MODEL CHECKS FAILED")'
    ws["A2"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    ws["A2"].font = Font(color="006100", bold=True)
    ws["A2"].alignment = Alignment(horizontal="left")
    banner_range = f"A2:{get_column_letter(last_col)}2"
    ws.conditional_formatting.add(
        banner_range,
        FormulaRule(formula=["Model_Checks!$B$2=0"], fill=PatternFill("solid", fgColor=LIGHT_GREEN), font=Font(color="006100", bold=True)),
    )
    ws.conditional_formatting.add(
        banner_range,
        FormulaRule(formula=["Model_Checks!$B$2>0"], fill=PatternFill("solid", fgColor=LIGHT_RED), font=Font(color="9C0006", bold=True)),
    )
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    ws["A3"] = "Amok Equipment AS | NOK unless stated | Generated 2026-10-10 | Version 1.0"
    ws["A3"].font = Font(italic=True, color=GRAY, size=9)


def section(ws, row: int, label: str, last_col: int = 9) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row, 1, label)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(color=WHITE, bold=True)


def header_row(ws, row: int, start_col: int, end_col: int, fill: str = DARK_BLUE) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_input(cell, value, number_format: str, comment: str | None = None) -> None:
    cell.value = value
    cell.number_format = number_format
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.font = Font(color=INPUT_BLUE)
    if comment:
        cell.comment = Comment(comment, "OpenAI")


def set_formula(cell, formula: str, number_format: str | None = None) -> None:
    cell.value = formula
    cell.font = Font(color=FORMULA_GREEN)
    if number_format:
        cell.number_format = number_format


def apply_timeline(ws, row: int = 4) -> None:
    ws.cell(row, 1, "Fiscal year")
    ws.cell(row, 1).font = Font(bold=True, color=WHITE)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=DARK_BLUE)
    for year, col in COLS.items():
        cell = ws.cell(row, col, year)
        cell.number_format = FMT_YEAR
        cell.alignment = Alignment(horizontal="center")
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE if year <= 2025 else BLUE)
        if year == 2026:
            cell.comment = Comment("Forecast period begins; all assumptions are editable on Assumptions_and_Inputs.", "OpenAI")


def set_year_row_formula(ws, row: int, year: int, formula: str, number_format: str = FMT_NOK) -> None:
    cell = ws.cell(row, COLS[year])
    set_formula(cell, formula, number_format)


def main() -> None:
    financials, cash_flows = read_historical_records()
    wb = Workbook()
    summary = wb.active
    summary.title = "TOC_and_Summary"
    assumptions = wb.create_sheet("Assumptions_and_Inputs")
    statements = wb.create_sheet("Financial_Statements")
    schedules = wb.create_sheet("Supporting_Schedules")
    scenario_analysis = wb.create_sheet("Scenario_Analysis")
    sensitivity = wb.create_sheet("Sensitivity_Analysis")
    checks = wb.create_sheet("Model_Checks")
    historical = wb.create_sheet("Historical_Data")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    for ws in wb.worksheets:
        style_sheet(ws)

    # Historical source inputs and their transparent mapping.
    title(historical, "Historical_Data | Reported source data and explicit source mappings", 5)
    historical["A4"] = "Source annual accounts (NOK)"
    for year, col in zip(ACTUAL_YEARS, range(2, 5)):
        historical.cell(4, col, year)
    header_row(historical, 4, 1, 4)
    historical_labels = {
        5: ("Revenue", "salgsinntekter"),
        6: ("Reported total operating expenses", "driftskostnader"),
        7: ("Personnel expenses", "lonnskostnad"),
        8: ("Reported operating profit (EBIT)", "driftsresultat"),
        9: ("Net financial income / (expense)", "netto_finans"),
        10: ("Reported income tax expense / (benefit)", "skattekostnad"),
        11: ("Reported net income", "arsresultat"),
        12: ("Net fixed assets", "anleggsmidler"),
        13: ("Total current assets", "omlopsmidler"),
        14: ("Inventory (reported; 2025 not separately disclosed)", "varer"),
        15: ("Accounts receivable", "fordringer"),
        16: ("Cash and cash equivalents", "bank"),
        17: ("Reported total assets", "sum_eiendeler"),
        18: ("Short-term liabilities", "kortsiktig_gjeld"),
        19: ("Long-term liabilities", "langsiktig_gjeld"),
        20: ("Reported total liabilities", "sum_gjeld"),
        21: ("Reported shareholders' equity", "egenkapital"),
    }
    for row, (label, key) in historical_labels.items():
        historical.cell(row, 1, label)
        for year, col in zip(ACTUAL_YEARS, range(2, 5)):
            value = financials[year].get(key)
            cell = historical.cell(row, col, value)
            if value is not None:
                cell.number_format = FMT_NOK
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
                cell.font = Font(color=INPUT_BLUE)
                cell.comment = Comment(f"Source: amok_equipment_alle_aar_wide.csv, {year}, field '{key}'.", "OpenAI")
            else:
                cell.value = None
    historical["A23"] = "Cash flow source data (NOK; reported cash-flow reconciliation inputs)"
    historical["A23"].fill = PatternFill("solid", fgColor=BLUE)
    historical["A23"].font = Font(color=WHITE, bold=True)
    for col in range(2, 5):
        historical.cell(24, col, ACTUAL_YEARS[col - 2])
    header_row(historical, 24, 1, 4)
    cf_rows = {
        25: ("Net income (cash flow source)", "Årsresultat"),
        26: ("Depreciation and amortization", "Avskrivninger"),
        27: ("Accounts receivable cash-flow impact", "Endring kundefordringer"),
        28: ("Inventory cash-flow impact", "Endring varelager"),
        29: ("Accounts payable cash-flow impact", "Endring leverandørgjeld"),
        30: ("Reported estimated cash from operations", "Estimert kontantstrøm fra drift"),
        31: ("Opening cash", "Åpningsbank"),
        32: ("Closing cash", "Sluttbank"),
        33: ("Unallocated other cash flows (source)", "Ufordelt øvrig kontantstrøm"),
    }
    for row, (label, source_label) in cf_rows.items():
        historical.cell(row, 1, label)
        for year, col in zip(ACTUAL_YEARS, range(2, 5)):
            value = cash_flows.get(year, {}).get(source_label)
            if row == 26 and year == 2023:
                value = 0
                note = "No 2023 D&A detail was supplied. Zero is a presentation proxy, not a reported figure."
            elif value is not None:
                note = f"Source: amok_cash_flow_2024_2025.csv, {year}, '{source_label}'."
            else:
                note = None
            cell = historical.cell(row, col, value)
            if value is not None:
                cell.number_format = FMT_NOK
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
                cell.font = Font(color=INPUT_BLUE)
                if note:
                    cell.comment = Comment(note, "OpenAI")
    historical["A35"] = "Source notes"
    historical["A35"].font = Font(bold=True, color=WHITE)
    historical["A35"].fill = PatternFill("solid", fgColor=BLUE)
    historical.merge_cells("A36:E36")
    historical["A36"] = "2025 inventory was not separately reported; the model treats current assets excluding cash and receivables as inventory / other operating current assets."
    historical.merge_cells("A37:E37")
    historical["A37"] = "2024 reported balance sheet totals differ by NOK 1 due to source rounding. Historical cash flows include unallocated cash movements because investing/financing detail was not supplied."
    historical.merge_cells("A38:E38")
    historical["A38"] = "Annual statements are historical reported figures. Financial_Statements derives non-personnel operating expenses as reported total operating expenses less reported personnel expenses."
    for row in range(36, 39):
        historical.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        historical.row_dimensions[row].height = 32

    # Centralized assumptions and scenario selector.
    title(assumptions, "Assumptions_and_Inputs | Centralized scenario drivers", 6)
    assumptions["A4"] = "Active scenario"
    set_input(
        assumptions["B4"],
        "Base Case",
        "General",
        "Select a complete, internally coherent planning scenario. Forecasts and sensitivities update on recalculation.",
    )
    assumptions["A5"] = "Driver (forecast assumption)"
    for index, scenario in enumerate(SCENARIOS, start=2):
        assumptions.cell(5, index, scenario)
    assumptions["F5"] = "Active assumption"
    header_row(assumptions, 5, 1, 6)
    scenario_rows = {
        6: ("Annual volume growth", [0.02, 0.10, -0.04, -0.15], FMT_PCT),
        7: ("Annual price realization growth", [0.02, 0.03, 0.01, -0.02], FMT_PCT),
        8: ("Non-personnel operating costs / revenue", [0.70, 0.68, 0.83, 0.90], FMT_PCT),
        9: ("Annual wage inflation", [0.035, 0.03, 0.06, 0.08], FMT_PCT),
        10: ("Accounts receivable days (DSO)", [43.0, 38.0, 50.0, 60.0], FMT_DAYS),
        11: ("Inventory / other operating assets days (DIO)", [325.0, 300.0, 360.0, 400.0], FMT_DAYS),
        12: ("Accounts payable days (DPO)", [65.0, 70.0, 60.0, 50.0], FMT_DAYS),
        13: ("Capital expenditure / revenue", [0.02, 0.025, 0.015, 0.01], FMT_PCT),
        14: ("Depreciation rate on opening net PP&E + half-year CapEx", [0.15, 0.15, 0.18, 0.20], FMT_PCT),
        15: ("Cash interest yield", [0.025, 0.03, 0.02, 0.015], FMT_PCT),
        16: ("Revolver interest rate", [0.07, 0.07, 0.09, 0.12], FMT_PCT),
        17: ("Cash tax rate (no loss-carryforward benefit)", [0.22, 0.22, 0.22, 0.22], FMT_PCT),
        18: ("Dividend payout ratio of positive net income", [0.0, 0.0, 0.0, 0.0], FMT_PCT),
        19: ("Minimum cash target (NOK)", [1_000_000, 1_000_000, 1_000_000, 1_000_000], FMT_NOK),
        20: ("Revolver commitment (NOK)", [5_000_000, 5_000_000, 5_000_000, 5_000_000], FMT_NOK),
        21: ("Excess-cash debt repayment sweep", [0.5, 0.5, 0.5, 0.5], FMT_PCT),
        22: ("Illustrative scenario probability", [0.55, 0.20, 0.20, 0.05], FMT_PCT),
    }
    for row, (label, values, number_format) in scenario_rows.items():
        assumptions.cell(row, 1, label)
        for col, value in enumerate(values, start=2):
            comment = "Illustrative forecast assumption only; replace with management-approved inputs." if row < 22 else "Illustrative probability weight; scenario weights must sum to 100%."
            set_input(assumptions.cell(row, col), value, number_format, comment)
        set_formula(assumptions.cell(row, 6), f"=IFERROR(INDEX($B{row}:$E{row},1,MATCH($B$4,$B$5:$E$5,0)),0)", number_format)
    assumptions["A24"] = "Days per year (working capital)"
    set_input(assumptions["B24"], 365, FMT_DAYS, "Annual day-count convention used in DSO, DIO, and DPO schedules.")
    assumptions["A25"] = "Model check tolerance (NOK)"
    set_input(assumptions["B25"], 1.01, FMT_NOK, "Tolerance used for balance-sheet and reconciliation checks.")
    assumptions["A26"] = "Forecast years"
    set_formula(assumptions["B26"], "=COUNT(Financial_Statements!$E$4:$I$4)", FMT_YEAR)
    assumptions["A27"] = "Model convention"
    assumptions.merge_cells("B27:F27")
    assumptions["B27"] = "Annual forecast; interest uses opening cash/debt to prevent circularity. No unmanaged circular references."
    assumptions["B27"].alignment = Alignment(wrap_text=True)
    assumptions["A28"] = "Actual / source convention"
    assumptions.merge_cells("B28:F28")
    assumptions["B28"] = "Blue cells are hardcoded inputs; green-font cells are formulas. Forecast drivers are scenarios, not company guidance."
    assumptions["B28"].alignment = Alignment(wrap_text=True)
    choices = DataValidation(type="list", formula1="=ScenarioList", allow_blank=False)
    choices.error = "Select a scenario from the list."
    choices.errorTitle = "Invalid scenario"
    choices.prompt = "Choose the case applied to the integrated forecast."
    choices.promptTitle = "Active scenario"
    assumptions.add_data_validation(choices)
    choices.add(assumptions["B4"])
    wb.defined_names.add(DefinedName("ScenarioList", attr_text="'Assumptions_and_Inputs'!$B$5:$E$5"))

    # Supporting schedules.
    title(schedules, "Supporting_Schedules | Working capital, PP&E, debt, and equity roll-forwards", 9)
    apply_timeline(schedules)
    sched_rows = {
        "wc_header": 6,
        "ar": 7,
        "inventory": 8,
        "ap": 9,
        "ar_change": 10,
        "inventory_change": 11,
        "ap_change": 12,
        "ppe_header": 14,
        "opening_ppe": 15,
        "capex": 16,
        "depreciation": 17,
        "closing_ppe": 18,
        "debt_header": 20,
        "opening_debt": 21,
        "draw": 22,
        "repayment": 23,
        "closing_debt": 24,
        "cash_interest": 25,
        "debt_interest": 26,
        "equity_header": 28,
        "opening_equity": 29,
        "net_income": 30,
        "dividends": 31,
        "closing_equity": 32,
    }
    for row, label in [
        (6, "Working capital schedule (NOK)"),
        (14, "PP&E corkscrew (NOK)"),
        (20, "Revolver and interest (NOK)"),
        (28, "Shareholders' equity corkscrew (NOK)"),
    ]:
        section(schedules, row, label)
    for row, label in [
        (7, "Accounts receivable"),
        (8, "Inventory and other operating current assets"),
        (9, "Accounts payable and other current liabilities (proxy)"),
        (10, "Change in accounts receivable (balance change)"),
        (11, "Change in inventory / other operating assets"),
        (12, "Change in accounts payable / current liabilities"),
        (15, "Opening net PP&E"),
        (16, "Capital expenditures"),
        (17, "Depreciation and amortization"),
        (18, "Closing net PP&E"),
        (21, "Opening revolver debt"),
        (22, "Revolver draw"),
        (23, "Revolver principal repayment"),
        (24, "Closing revolver debt"),
        (25, "Cash interest income (opening cash method)"),
        (26, "Revolver interest expense (opening debt method)"),
        (29, "Opening shareholders' equity"),
        (30, "Net income"),
        (31, "Dividends"),
        (32, "Closing shareholders' equity"),
    ]:
        schedules.cell(row, 1, label)

    # Integrated income statement, balance sheet, and cash flow statement.
    title(statements, "Financial_Statements | Integrated income statement, balance sheet, and cash flow", 9)
    apply_timeline(statements)
    statement_rows = {
        5: "Income statement (NOK, except percentages)",
        23: "Balance sheet (NOK)",
        38: "Cash flow statement (NOK)",
    }
    for row, label in statement_rows.items():
        section(statements, row, label)
    statement_labels = {
        6: "Revenue",
        7: "Revenue growth",
        8: "Non-personnel operating expenses (proxy)",
        9: "Pre-personnel operating contribution",
        10: "Pre-personnel contribution margin",
        11: "Personnel expenses",
        12: "EBITDA",
        13: "EBITDA margin",
        14: "Depreciation and amortization",
        15: "Operating profit (EBIT)",
        16: "EBIT margin",
        17: "Net interest income / (expense)",
        18: "Earnings before tax",
        19: "Income tax expense / (benefit)",
        20: "Net income",
        21: "Net income margin",
        24: "Net PP&E",
        25: "Cash and cash equivalents",
        26: "Accounts receivable",
        27: "Inventory and other operating current assets",
        28: "Total current assets",
        29: "Total assets",
        30: "Accounts payable / other current liabilities",
        31: "Long-term debt / revolver",
        32: "Total liabilities",
        33: "Shareholders' equity",
        34: "Total liabilities and equity",
        35: "Balance sheet check (assets less claims)",
        39: "Net income",
        40: "Depreciation and amortization",
        41: "Accounts receivable cash-flow impact",
        42: "Inventory / other operating assets cash-flow impact",
        43: "Accounts payable / current liabilities cash-flow impact",
        44: "Net working capital cash-flow impact",
        45: "Cash from operations",
        46: "Capital expenditure",
        47: "Revolver draw",
        48: "Revolver principal repayment",
        49: "Dividends",
        50: "Unallocated other cash flows (historical source only)",
        51: "Cash from investing",
        52: "Cash from financing",
        53: "Net change in cash",
        54: "Opening cash",
        55: "Ending cash",
    }
    for row, label in statement_labels.items():
        statements.cell(row, 1, label)

    # Historical schedules and statements.
    for year in ACTUAL_YEARS:
        col = COLS[year]
        data_col = year - 2021
        hist_col = chr(64 + data_col)
        source_ref = lambda row: f"Historical_Data!${hist_col}${row}"
        set_year_row_formula(statements, 6, year, f"={source_ref(5)}")
        if year == 2023:
            set_year_row_formula(statements, 7, year, '=""', FMT_PCT)
        else:
            prev_col = chr(col + 63)
            set_year_row_formula(statements, 7, year, f'=IFERROR({hist_col}6/{prev_col}6-1,0)', FMT_PCT)
        set_year_row_formula(statements, 8, year, f"={source_ref(6)}-{source_ref(7)}")
        set_year_row_formula(statements, 9, year, f"={hist_col}6-{hist_col}8")
        set_year_row_formula(statements, 10, year, f"=IFERROR({hist_col}9/{hist_col}6,0)", FMT_PCT)
        set_year_row_formula(statements, 11, year, f"={source_ref(7)}")
        set_year_row_formula(statements, 14, year, f"=Historical_Data!{hist_col}26")
        set_year_row_formula(statements, 15, year, f"={source_ref(8)}")
        set_year_row_formula(statements, 12, year, f"={hist_col}15+{hist_col}14")
        set_year_row_formula(statements, 13, year, f"=IFERROR({hist_col}12/{hist_col}6,0)", FMT_PCT)
        set_year_row_formula(statements, 16, year, f"=IFERROR({hist_col}15/{hist_col}6,0)", FMT_PCT)
        set_year_row_formula(statements, 17, year, f"={source_ref(9)}")
        set_year_row_formula(statements, 18, year, f"={hist_col}15+{hist_col}17")
        set_year_row_formula(statements, 19, year, f"={source_ref(10)}")
        set_year_row_formula(statements, 20, year, f"={source_ref(11)}")
        set_year_row_formula(statements, 21, year, f"=IFERROR({hist_col}20/{hist_col}6,0)", FMT_PCT)

        set_year_row_formula(schedules, 7, year, f"=Historical_Data!{hist_col}15")
        set_year_row_formula(schedules, 8, year, f"=Historical_Data!{hist_col}13-Historical_Data!{hist_col}15-Historical_Data!{hist_col}16")
        set_year_row_formula(schedules, 9, year, f"=Historical_Data!{hist_col}18")
        if year == 2023:
            for row in (10, 11, 12):
                set_year_row_formula(schedules, row, year, '=""')
        else:
            prev_col = chr(col + 63)
            set_year_row_formula(schedules, 10, year, f"={hist_col}7-{prev_col}7")
            set_year_row_formula(schedules, 11, year, f"={hist_col}8-{prev_col}8")
            set_year_row_formula(schedules, 12, year, f"={hist_col}9-{prev_col}9")
        set_year_row_formula(schedules, 15, year, '=""')
        set_year_row_formula(schedules, 16, year, '=""')
        set_year_row_formula(schedules, 17, year, f"=Historical_Data!{hist_col}26")
        set_year_row_formula(schedules, 18, year, f"=Historical_Data!{hist_col}12")
        for row in (21, 22, 23, 25, 26, 29, 31):
            set_year_row_formula(schedules, row, year, '=""')
        set_year_row_formula(schedules, 24, year, f"=Historical_Data!{hist_col}19")
        set_year_row_formula(schedules, 30, year, f"=Historical_Data!{hist_col}11")
        set_year_row_formula(schedules, 32, year, f"=Historical_Data!{hist_col}21")

        for row, formula in {
            24: f"=Historical_Data!{hist_col}12",
            25: f"=Historical_Data!{hist_col}16",
            26: f"=Historical_Data!{hist_col}15",
            27: f"=Historical_Data!{hist_col}13-Historical_Data!{hist_col}15-Historical_Data!{hist_col}16",
            28: f"=SUM({hist_col}25:{hist_col}27)",
            29: f"=Historical_Data!{hist_col}17",
            30: f"=Historical_Data!{hist_col}18",
            31: f"=Historical_Data!{hist_col}19",
            32: f"=Historical_Data!{hist_col}20",
            33: f"=Historical_Data!{hist_col}21",
            34: f"=SUM({hist_col}32:{hist_col}33)",
            35: f"={hist_col}29-{hist_col}34",
        }.items():
            set_year_row_formula(statements, row, year, formula)

        if year == 2023:
            for row in range(39, 56):
                set_year_row_formula(statements, row, year, '=""')
        else:
            cf_map = {39: 25, 40: 26, 41: 27, 42: 28, 43: 29, 50: 33, 54: 31}
            for row, source_row in cf_map.items():
                set_year_row_formula(statements, row, year, f"=Historical_Data!{hist_col}{source_row}")
            set_year_row_formula(statements, 44, year, f"=SUM({hist_col}41:{hist_col}43)")
            set_year_row_formula(statements, 45, year, f"={hist_col}39+{hist_col}40+{hist_col}44")
            for row in (46, 47, 48, 49, 51, 52):
                set_year_row_formula(statements, row, year, "=0")
            set_year_row_formula(statements, 53, year, f"=SUM({hist_col}45,{hist_col}50:{hist_col}52)")
            set_year_row_formula(statements, 55, year, f"={hist_col}54+{hist_col}53")

    # Forecast period: active scenario assumptions link all three statements.
    for year in FORECAST_YEARS:
        col = COLS[year]
        prev_col = chr(col + 63)
        letter = chr(col + 64)
        assumption = "Assumptions_and_Inputs"
        set_year_row_formula(statements, 6, year, f"={prev_col}6*(1+{assumption}!$F$6)*(1+{assumption}!$F$7)")
        set_year_row_formula(statements, 7, year, f"=IFERROR({letter}6/{prev_col}6-1,0)", FMT_PCT)
        set_year_row_formula(statements, 8, year, f"={letter}6*{assumption}!$F$8")
        set_year_row_formula(statements, 9, year, f"={letter}6-{letter}8")
        set_year_row_formula(statements, 10, year, f"=IFERROR({letter}9/{letter}6,0)", FMT_PCT)
        set_year_row_formula(statements, 11, year, f"={prev_col}11*(1+{assumption}!$F$6)*(1+{assumption}!$F$9)")
        set_year_row_formula(statements, 12, year, f"={letter}9-{letter}11")
        set_year_row_formula(statements, 13, year, f"=IFERROR({letter}12/{letter}6,0)", FMT_PCT)
        set_year_row_formula(statements, 14, year, f"=Supporting_Schedules!{letter}17")
        set_year_row_formula(statements, 15, year, f"={letter}12-{letter}14")
        set_year_row_formula(statements, 16, year, f"=IFERROR({letter}15/{letter}6,0)", FMT_PCT)
        set_year_row_formula(statements, 17, year, f"=Supporting_Schedules!{letter}25-Supporting_Schedules!{letter}26")
        set_year_row_formula(statements, 18, year, f"={letter}15+{letter}17")
        set_year_row_formula(statements, 19, year, f"=MAX(0,{letter}18)*{assumption}!$F$17")
        set_year_row_formula(statements, 20, year, f"={letter}18-{letter}19")
        set_year_row_formula(statements, 21, year, f"=IFERROR({letter}20/{letter}6,0)", FMT_PCT)

        set_year_row_formula(schedules, 7, year, f"=IFERROR(Financial_Statements!{letter}6/{assumption}!$B$24*{assumption}!$F$10,0)")
        set_year_row_formula(schedules, 8, year, f"=IFERROR(Financial_Statements!{letter}8/{assumption}!$B$24*{assumption}!$F$11,0)")
        set_year_row_formula(schedules, 9, year, f"=IFERROR(Financial_Statements!{letter}8/{assumption}!$B$24*{assumption}!$F$12,0)")
        set_year_row_formula(schedules, 10, year, f"={letter}7-{prev_col}7")
        set_year_row_formula(schedules, 11, year, f"={letter}8-{prev_col}8")
        set_year_row_formula(schedules, 12, year, f"={letter}9-{prev_col}9")
        set_year_row_formula(schedules, 15, year, f"={prev_col}18")
        set_year_row_formula(schedules, 16, year, f"=Financial_Statements!{letter}6*{assumption}!$F$13")
        set_year_row_formula(schedules, 17, year, f"=MIN({letter}15+{letter}16,({letter}15+{letter}16/2)*{assumption}!$F$14)")
        set_year_row_formula(schedules, 18, year, f"=MAX(0,{letter}15+{letter}16-{letter}17)")

        set_year_row_formula(schedules, 21, year, f"={prev_col}24")
        set_year_row_formula(schedules, 25, year, f"=Financial_Statements!{prev_col}25*{assumption}!$F$15")
        set_year_row_formula(schedules, 26, year, f"={letter}21*{assumption}!$F$16")
        set_year_row_formula(schedules, 22, year, f"=MIN(MAX(0,{assumption}!$F$20-{letter}21),MAX(0,{assumption}!$F$19-(Financial_Statements!{prev_col}25+Financial_Statements!{letter}45+Financial_Statements!{letter}46+Financial_Statements!{letter}49)))")
        set_year_row_formula(schedules, 23, year, f"=MIN({letter}21+{letter}22,MAX(0,Financial_Statements!{prev_col}25+Financial_Statements!{letter}45+Financial_Statements!{letter}46+Financial_Statements!{letter}49-{assumption}!$F$19)*{assumption}!$F$21)")
        set_year_row_formula(schedules, 24, year, f"={letter}21+{letter}22-{letter}23")

        set_year_row_formula(schedules, 29, year, f"={prev_col}32")
        set_year_row_formula(schedules, 30, year, f"=Financial_Statements!{letter}20")
        set_year_row_formula(schedules, 31, year, f"=MAX(0,{letter}30)*{assumption}!$F$18")
        set_year_row_formula(schedules, 32, year, f"={letter}29+{letter}30-{letter}31")

        for row, formula in {
            24: f"=Supporting_Schedules!{letter}18",
            25: f"={letter}55",
            26: f"=Supporting_Schedules!{letter}7",
            27: f"=Supporting_Schedules!{letter}8",
            28: f"=SUM({letter}25:{letter}27)",
            29: f"={letter}24+{letter}28",
            30: f"=Supporting_Schedules!{letter}9",
            31: f"=Supporting_Schedules!{letter}24",
            32: f"=SUM({letter}30:{letter}31)",
            33: f"=Supporting_Schedules!{letter}32",
            34: f"=SUM({letter}32:{letter}33)",
            35: f"={letter}29-{letter}34",
            39: f"={letter}20",
            40: f"={letter}14",
            41: f"=-Supporting_Schedules!{letter}10",
            42: f"=-Supporting_Schedules!{letter}11",
            43: f"=Supporting_Schedules!{letter}12",
            44: f"=SUM({letter}41:{letter}43)",
            45: f"={letter}39+{letter}40+{letter}44",
            46: f"=-Supporting_Schedules!{letter}16",
            47: f"=Supporting_Schedules!{letter}22",
            48: f"=-Supporting_Schedules!{letter}23",
            49: f"=-Supporting_Schedules!{letter}31",
            50: "=0",
            51: f"={letter}46",
            52: f"=SUM({letter}47:{letter}49)",
            53: f"=SUM({letter}45,{letter}50:{letter}52)",
            54: f"={prev_col}55",
            55: f"={letter}54+{letter}53",
        }.items():
            set_year_row_formula(statements, row, year, formula)

    # Income statement historical links use source-based expense reconstruction.
    statements["A57"] = "Historical mapping"
    statements["A57"].font = Font(bold=True, color=WHITE)
    statements["A57"].fill = PatternFill("solid", fgColor=BLUE)
    statements.merge_cells("A58:I58")
    statements["A58"] = "The source does not separately disclose cost of sales. Non-personnel operating expenses are the residual of reported total operating expenses less personnel expenses; this is a proxy, not a reported COGS figure."
    statements["A58"].alignment = Alignment(wrap_text=True)
    statements.merge_cells("A59:I59")
    statements["A59"] = "The 2025 inventory field is blank in the source; the model presents current assets excluding cash and receivables as inventory / other operating current assets. Historical CFO uses supplied cash-flow working-capital effects."
    statements["A59"].alignment = Alignment(wrap_text=True)

    # Scenario comparison: explicit 2030 revenue and EBITDA for each coherent case.
    title(scenario_analysis, "Scenario_Analysis | Side-by-side outcomes and forecast schedules by scenario", 6)
    scenario_analysis["A4"] = "2030 output (NOK, except margin)"
    for col, scenario in enumerate(SCENARIOS, start=2):
        scenario_analysis.cell(4, col, scenario)
    header_row(scenario_analysis, 4, 1, 5)
    scenario_analysis["A5"] = "Revenue"
    scenario_analysis["A6"] = "EBITDA"
    scenario_analysis["A7"] = "EBITDA margin"
    scenario_analysis["A8"] = "Ending cash"
    scenario_analysis["A9"] = "Closing revolver debt"
    scenario_analysis["A10"] = "Liquidity status"
    n_years = "Assumptions_and_Inputs!$B$26"
    case_blocks: dict[str, dict[str, int]] = {}
    case_metrics = [
        (2, "Revenue"),
        (3, "Non-personnel operating costs"),
        (4, "Personnel expenses"),
        (5, "EBITDA"),
        (6, "Depreciation and amortization"),
        (7, "EBIT"),
        (8, "Cash interest income"),
        (9, "Revolver interest expense"),
        (10, "Net income"),
        (11, "Accounts receivable"),
        (12, "Inventory and other operating assets"),
        (13, "Accounts payable / current liabilities"),
        (14, "Capital expenditure"),
        (15, "Opening net PP&E"),
        (16, "Closing net PP&E"),
        (17, "Opening cash"),
        (18, "Cash from operations"),
        (19, "Dividends"),
        (20, "Cash before financing"),
        (21, "Opening revolver debt"),
        (22, "Revolver draw"),
        (23, "Revolver repayment"),
        (24, "Closing revolver debt"),
        (25, "Ending cash"),
        (26, "Opening shareholders' equity"),
        (27, "Closing shareholders' equity"),
        (28, "Balance sheet difference"),
        (29, "Liquidity flag (1 = breach)"),
    ]
    for col, input_col in zip(range(2, 6), ["B", "C", "D", "E"]):
        growth = f"Assumptions_and_Inputs!${input_col}$6"
        price = f"Assumptions_and_Inputs!${input_col}$7"
        cost_ratio = f"Assumptions_and_Inputs!${input_col}$8"
        wage = f"Assumptions_and_Inputs!${input_col}$9"
        block = 13 + (col - 2) * 30
        case_blocks[SCENARIOS[col - 2]] = {
            "base": block,
            "input_col": ord(input_col) - 64,
            "cash": block + 25,
            "debt": block + 24,
            "balance": block + 28,
            "liquidity": block + 29,
        }
        section(scenario_analysis, block, f"{SCENARIOS[col - 2]} | Five-year integrated scenario schedule", 6)
        scenario_analysis.cell(block + 1, 1, "Forecast year")
        for year_index, year in enumerate(FORECAST_YEARS, start=2):
            scenario_analysis.cell(block + 1, year_index, year)
        header_row(scenario_analysis, block + 1, 1, 6)
        for offset, label in case_metrics:
            scenario_analysis.cell(block + offset, 1, label)
        for forecast_index, model_col in enumerate(range(2, 7)):
            scenario_col = get_column_letter(model_col)
            prior_col = get_column_letter(model_col - 1) if model_col > 2 else None
            prior_rev = f"{prior_col}{block+2}" if prior_col else "Historical_Data!$D$5"
            prior_pay = f"{prior_col}{block+4}" if prior_col else "Historical_Data!$D$7"
            prior_ppe = f"{prior_col}{block+16}" if prior_col else "Historical_Data!$D$12"
            prior_cash = f"{prior_col}{block+25}" if prior_col else "Historical_Data!$D$16"
            prior_debt = f"{prior_col}{block+24}" if prior_col else "IFERROR(Historical_Data!$D$19,0)"
            prior_equity = f"{prior_col}{block+27}" if prior_col else "Historical_Data!$D$21"
            prior_ar = f"{prior_col}{block+11}" if prior_col else "Historical_Data!$D$15"
            prior_inventory = f"{prior_col}{block+12}" if prior_col else "(Historical_Data!$D$13-Historical_Data!$D$15-Historical_Data!$D$16)"
            prior_ap = f"{prior_col}{block+13}" if prior_col else "Historical_Data!$D$18"
            formulas_by_offset = {
                2: f"={prior_rev}*(1+{growth})*(1+{price})",
                3: f"={scenario_col}{block+2}*{cost_ratio}",
                4: f"={prior_pay}*(1+{growth})*(1+{wage})",
                5: f"={scenario_col}{block+2}-{scenario_col}{block+3}-{scenario_col}{block+4}",
                6: f"=MIN({scenario_col}{block+15}+{scenario_col}{block+14},({scenario_col}{block+15}+{scenario_col}{block+14}/2)*Assumptions_and_Inputs!${input_col}$14)",
                7: f"={scenario_col}{block+5}-{scenario_col}{block+6}",
                8: f"={scenario_col}{block+17}*Assumptions_and_Inputs!${input_col}$15",
                9: f"={scenario_col}{block+21}*Assumptions_and_Inputs!${input_col}$16",
                10: f"={scenario_col}{block+7}+{scenario_col}{block+8}-{scenario_col}{block+9}-MAX(0,{scenario_col}{block+7}+{scenario_col}{block+8}-{scenario_col}{block+9})*Assumptions_and_Inputs!${input_col}$17",
                11: f"=IFERROR({scenario_col}{block+2}/Assumptions_and_Inputs!$B$24*Assumptions_and_Inputs!${input_col}$10,0)",
                12: f"=IFERROR({scenario_col}{block+3}/Assumptions_and_Inputs!$B$24*Assumptions_and_Inputs!${input_col}$11,0)",
                13: f"=IFERROR({scenario_col}{block+3}/Assumptions_and_Inputs!$B$24*Assumptions_and_Inputs!${input_col}$12,0)",
                14: f"={scenario_col}{block+2}*Assumptions_and_Inputs!${input_col}$13",
                15: f"={prior_ppe}",
                16: f"=MAX(0,{scenario_col}{block+15}+{scenario_col}{block+14}-{scenario_col}{block+6})",
                17: f"={prior_cash}",
                18: (
                    f"={scenario_col}{block+10}+{scenario_col}{block+6}"
                    f"-({scenario_col}{block+11}-{prior_ar})"
                    f"-({scenario_col}{block+12}-{prior_inventory})"
                    f"+({scenario_col}{block+13}-{prior_ap})"
                ),
                19: f"=MAX(0,{scenario_col}{block+10})*Assumptions_and_Inputs!${input_col}$18",
                20: f"={scenario_col}{block+17}+{scenario_col}{block+18}-{scenario_col}{block+14}-{scenario_col}{block+19}",
                21: f"={prior_debt}",
                22: f"=MIN(MAX(0,Assumptions_and_Inputs!${input_col}$20-{scenario_col}{block+21}),MAX(0,Assumptions_and_Inputs!${input_col}$19-{scenario_col}{block+20}))",
                23: f"=MIN({scenario_col}{block+21}+{scenario_col}{block+22},MAX(0,{scenario_col}{block+20}-Assumptions_and_Inputs!${input_col}$19)*Assumptions_and_Inputs!${input_col}$21)",
                24: f"={scenario_col}{block+21}+{scenario_col}{block+22}-{scenario_col}{block+23}",
                25: f"={scenario_col}{block+20}+{scenario_col}{block+22}-{scenario_col}{block+23}",
                26: f"={prior_equity}",
                27: f"={scenario_col}{block+26}+{scenario_col}{block+10}-{scenario_col}{block+19}",
                28: f"={scenario_col}{block+25}+{scenario_col}{block+11}+{scenario_col}{block+12}+{scenario_col}{block+16}-{scenario_col}{block+13}-{scenario_col}{block+24}-{scenario_col}{block+27}",
                29: f"=IF(OR({scenario_col}{block+25}<0,{scenario_col}{block+24}>Assumptions_and_Inputs!${input_col}$20),1,0)",
            }
            for offset, formula in formulas_by_offset.items():
                set_formula(scenario_analysis.cell(block + offset, model_col), formula, FMT_NOK)
        last_col = "F"
        set_formula(scenario_analysis.cell(5, col), f"={last_col}{block+2}", FMT_NOK)
        set_formula(scenario_analysis.cell(6, col), f"={last_col}{block+5}", FMT_NOK)
        set_formula(scenario_analysis.cell(7, col), f"=IFERROR({scenario_analysis.cell(6, col).coordinate}/{scenario_analysis.cell(5, col).coordinate},0)", FMT_PCT)
        set_formula(scenario_analysis.cell(8, col), f"={last_col}{block+25}", FMT_NOK)
        set_formula(scenario_analysis.cell(9, col), f"={last_col}{block+24}", FMT_NOK)
        set_formula(
            scenario_analysis.cell(10, col),
            f'=IF(OR({last_col}{block+25}<0,{last_col}{block+24}>Assumptions_and_Inputs!${input_col}$20),"BREACH",IF({last_col}{block+25}<Assumptions_and_Inputs!${input_col}$19-Assumptions_and_Inputs!$B$25,"BELOW MIN CASH","PASS"))',
        )
    header_row(scenario_analysis, 4, 1, 5)
    scenario_analysis.merge_cells("A11:F11")
    scenario_analysis["A11"] = "Each scenario schedule independently rolls working capital, PP&E, cash, revolver debt, and equity. Liquidity status flags cash shortfalls; active-case details also appear in Financial_Statements."
    scenario_analysis["A11"].alignment = Alignment(wrap_text=True, vertical="top")
    scenario_analysis.row_dimensions[11].height = 36

    # Formula-based two-way EBITDA sensitivity and one-way tornado table.
    title(sensitivity, "Sensitivity_Analysis | 2030 EBITDA scenario and driver sensitivity", 8)
    sensitivity["A4"] = "Two-way sensitivity | 2030 EBITDA (NOK)"
    sensitivity.merge_cells("A4:G4")
    sensitivity["A4"].fill = PatternFill("solid", fgColor=BLUE)
    sensitivity["A4"].font = Font(color=WHITE, bold=True)
    sensitivity["A5"] = "Cost ratio shocks (editable)"
    sensitivity["B5"] = "Input offset"
    ratio_offsets = [-0.05, -0.025, 0, 0.025, 0.05]
    volume_offsets = [-0.10, -0.05, 0, 0.05, 0.10]
    sensitivity["A6"] = "Volume growth Δ / Non-personnel cost ratio"
    for offset, col in zip(ratio_offsets, range(3, 8)):
        set_input(sensitivity.cell(5, col), offset, FMT_PCT, "Sensitivity shock to the selected scenario's non-personnel operating cost ratio.")
        set_formula(sensitivity.cell(6, col), f"=Assumptions_and_Inputs!$F$8+{get_column_letter(col)}$5", FMT_PCT)
    for offset, row in zip(volume_offsets, range(7, 12)):
        sensitivity.cell(row, 2, offset).number_format = FMT_PCT
        sensitivity.cell(row, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        sensitivity.cell(row, 2).font = Font(color=INPUT_BLUE)
        sensitivity.cell(row, 2).comment = Comment("Sensitivity shift in annual volume growth relative to active scenario.", "OpenAI")
        for col in range(3, 8):
            header = f"{chr(64 + col)}$6"
            volume_factor = f"(1+Assumptions_and_Inputs!$F$6+$B{row})"
            price_factor = "(1+Assumptions_and_Inputs!$F$7)"
            wage_factor = "(1+Assumptions_and_Inputs!$F$9)"
            formula = (
                f"=IFERROR(Historical_Data!$D$5*{volume_factor}^{n_years}*{price_factor}^{n_years}*(1-{header})"
                f"-Historical_Data!$D$7*{volume_factor}^{n_years}*{wage_factor}^{n_years},0)"
            )
            set_formula(sensitivity.cell(row, col), formula, FMT_NOK)
    header_row(sensitivity, 6, 2, 7)
    sensitivity.conditional_formatting.add(
        "C7:G11",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    sensitivity["A13"] = "Base case tie-out"
    set_formula(sensitivity["B13"], "=IFERROR(Financial_Statements!$I$12,0)", FMT_NOK)
    sensitivity["C13"] = "Sensitivity center"
    set_formula(sensitivity["D13"], "=E9", FMT_NOK)
    sensitivity["E13"] = "Difference"
    set_formula(sensitivity["F13"], "=D13-B13", FMT_NOK)

    sensitivity["A15"] = "One-way tornado | Change in 2030 EBITDA (NOK)"
    sensitivity.merge_cells("A15:F15")
    sensitivity["A15"].fill = PatternFill("solid", fgColor=BLUE)
    sensitivity["A15"].font = Font(color=WHITE, bold=True)
    for col, label in enumerate(["Driver", "Downside EBITDA", "Upside EBITDA", "Downside impact", "Upside impact", "Elasticity"], start=1):
        sensitivity.cell(16, col, label)
    header_row(sensitivity, 16, 1, 6)
    sensitivity["G16"] = "Downside shock"
    sensitivity["H16"] = "Upside shock"
    header_row(sensitivity, 16, 7, 8)
    tornado = [
        ("Volume growth", "growth", -0.05, 0.05),
        ("Price realization growth", "price", -0.02, 0.02),
        ("Non-personnel cost ratio", "cost", 0.03, -0.03),
        ("Wage inflation", "wage", 0.02, -0.02),
    ]

    def tornado_formula(variable: str, shock_cell: str, row: int) -> str:
        g = f"(Assumptions_and_Inputs!$F$6+{shock_cell})" if variable == "growth" else "Assumptions_and_Inputs!$F$6"
        p = f"(Assumptions_and_Inputs!$F$7+{shock_cell})" if variable == "price" else "Assumptions_and_Inputs!$F$7"
        c = f"(Assumptions_and_Inputs!$F$8+{shock_cell})" if variable == "cost" else "Assumptions_and_Inputs!$F$8"
        w = f"(Assumptions_and_Inputs!$F$9+{shock_cell})" if variable == "wage" else "Assumptions_and_Inputs!$F$9"
        return (
            f"=IFERROR(Historical_Data!$D$5*(1+{g})^{n_years}*(1+{p})^{n_years}*(1-{c})"
            f"-Historical_Data!$D$7*(1+{g})^{n_years}*(1+{w})^{n_years},0)"
        )

    for row, (label, variable, downside_shock, upside_shock) in enumerate(tornado, start=17):
        sensitivity.cell(row, 1, label)
        set_input(sensitivity.cell(row, 7), downside_shock, FMT_PCT, "Editable one-way downside shock for tornado sensitivity.")
        set_input(sensitivity.cell(row, 8), upside_shock, FMT_PCT, "Editable one-way upside shock for tornado sensitivity.")
        set_formula(sensitivity.cell(row, 2), tornado_formula(variable, f"$G{row}", row), FMT_NOK)
        set_formula(sensitivity.cell(row, 3), tornado_formula(variable, f"$H{row}", row), FMT_NOK)
        set_formula(sensitivity.cell(row, 4), f"=B{row}-$B$13", FMT_NOK)
        set_formula(sensitivity.cell(row, 5), f"=C{row}-$B$13", FMT_NOK)
        assumption_row = {"growth": 6, "price": 7, "cost": 8, "wage": 9}[variable]
        set_formula(
            sensitivity.cell(row, 6),
            f'=IFERROR(((C{row}-B{row})/$B$13)/(($H{row}-$G{row})/Assumptions_and_Inputs!$F${assumption_row}),0)',
            FMT_MULT,
        )
    sensitivity.merge_cells("A23:F23")
    sensitivity["A23"] = "Tornado shocks are one-at-a-time changes around the active scenario. Elasticity is the percentage change in EBITDA divided by the percentage change in the driver across the full downside-to-upside range."
    sensitivity["A23"].alignment = Alignment(wrap_text=True, vertical="top")
    sensitivity.row_dimensions[23].height = 34

    # Centralized model checks with historical and projected periods.
    title(checks, "Model_Checks | Accounting, roll-forward, scenario, and sensitivity assertions", 1)
    checks.merge_cells("A1:K1")
    checks.merge_cells("A3:K3")
    checks["A4"] = "Failed checks (zero is pass)"
    set_formula(checks["B4"], "=SUM(K6:K17)", FMT_NOK)
    checks["A5"] = "Check"
    checks["B5"] = "Test"
    for year, col in COLS.items():
        checks.cell(5, col + 1, year)  # C:J
    checks["K5"] = "Total failures"
    header_row(checks, 5, 1, 11)
    check_specs = [
        (6, "Balance sheet equilibrium", "Assets equal liabilities plus equity within the centralized NOK tolerance."),
        (7, "Cash reconciliation", "Cash flow ending cash agrees with balance sheet cash."),
        (8, "Equity roll-forward", "Forecast equity equals opening equity + net income - dividends."),
        (9, "PP&E roll-forward", "Forecast PP&E equals opening PP&E + CapEx - D&A."),
        (10, "Debt roll-forward / facility", "Debt rolls forward correctly and remains within facility; cash is not negative."),
        (11, "Core formula integrity", "Key statements are numeric; derived historical CFO ties to source."),
        (12, "Scenario probabilities", "Scenario probability weights sum to 100%."),
        (13, "Scenario ranges", "Growth, operating cost, and working-capital assumptions are within valid ranges."),
        (14, "Sensitivity integrity", "Two-way sensitivity grid is numeric and its center reconciles to 2030 EBITDA."),
        (15, "Scenario output integrity", "All four scenario revenue, EBITDA, cash, and debt outcomes are numeric."),
        (16, "Scenario balance-sheet integrity", "All scenario-specific balance sheets remain in accounting balance."),
        (17, "Selected-case reconciliation", "Selected scenario outputs tie to the integrated financial statements."),
    ]
    for row, label, description in check_specs:
        checks.cell(row, 1, label)
        checks.cell(row, 2, description)
    tolerance = "Assumptions_and_Inputs!$B$25"
    for year, col in COLS.items():
        model_col = chr(COLS[year] + 64)
        check_col = col + 1
        if year <= 2025:
            formulas = {
                6: f'=IF(ABS(Financial_Statements!{model_col}35)<={tolerance},0,1)',
                7: '=0' if year == 2023 else f'=IF(ABS(Financial_Statements!{model_col}55-Financial_Statements!{model_col}25)<={tolerance},0,1)',
                8: "=0",
                9: "=0",
                10: "=0",
                11: f'=IF(OR(ISERROR(Financial_Statements!{model_col}6),ISERROR(Financial_Statements!{model_col}29),ISERROR(Financial_Statements!{model_col}34){"" if year == 2023 else f",ABS(Financial_Statements!{model_col}45-Historical_Data!{model_col}30)>{tolerance}"}),1,0)',
                12: "=0",
                13: "=0",
                14: "=0",
                15: "=0",
                16: "=0",
                17: "=0",
            }
        else:
            formulas = {
                6: f'=IF(ABS(Financial_Statements!{model_col}35)<={tolerance},0,1)',
                7: f'=IF(ABS(Financial_Statements!{model_col}55-Financial_Statements!{model_col}25)<={tolerance},0,1)',
                8: f'=IF(ABS(Supporting_Schedules!{model_col}32-(Supporting_Schedules!{model_col}29+Supporting_Schedules!{model_col}30-Supporting_Schedules!{model_col}31))<={tolerance},0,1)',
                9: f'=IF(ABS(Supporting_Schedules!{model_col}18-(Supporting_Schedules!{model_col}15+Supporting_Schedules!{model_col}16-Supporting_Schedules!{model_col}17))<={tolerance},0,1)',
                10: f'=IF(OR(ABS(Supporting_Schedules!{model_col}24-(Supporting_Schedules!{model_col}21+Supporting_Schedules!{model_col}22-Supporting_Schedules!{model_col}23))>{tolerance},Supporting_Schedules!{model_col}24>Assumptions_and_Inputs!$F$20+{tolerance},Financial_Statements!{model_col}25<0,AND(Financial_Statements!{model_col}25<Assumptions_and_Inputs!$F$19-{tolerance},Supporting_Schedules!{model_col}24<Assumptions_and_Inputs!$F$20-{tolerance})),1,0)',
                11: f'=IF(OR(ISERROR(Financial_Statements!{model_col}6),ISERROR(Financial_Statements!{model_col}12),ISERROR(Financial_Statements!{model_col}20),ISERROR(Financial_Statements!{model_col}29),ISERROR(Financial_Statements!{model_col}34),ISERROR(Financial_Statements!{model_col}55)),1,0)',
                12: "=0",
                13: "=0",
                14: "=0",
                15: "=0",
                16: "=0",
                17: "=0",
            }
        for row, formula in formulas.items():
            set_formula(checks.cell(row, check_col), formula, FMT_NOK)
    checks["K6"] = "=SUM(C6:J6)"
    checks["K7"] = "=SUM(C7:J7)"
    checks["K8"] = "=SUM(C8:J8)"
    checks["K9"] = "=SUM(C9:J9)"
    checks["K10"] = "=SUM(C10:J10)"
    checks["K11"] = "=SUM(C11:J11)"
    set_formula(checks["K12"], '=IF(ABS(SUM(Assumptions_and_Inputs!$B$22:$E$22)-1)<=0.000001,0,1)', FMT_NOK)
    set_formula(checks["K13"], '=IF(OR(MIN(Assumptions_and_Inputs!$B$6:$E$6)<=-1,MIN(Assumptions_and_Inputs!$B$8:$E$8)<0,MAX(Assumptions_and_Inputs!$B$8:$E$8)>1,MIN(Assumptions_and_Inputs!$B$10:$E$12)<0),1,0)', FMT_NOK)
    set_formula(checks["K14"], f'=IF(OR(COUNT(Sensitivity_Analysis!$C$7:$G$11)<>25,ABS(Sensitivity_Analysis!$F$13)>{tolerance}),1,0)', FMT_NOK)
    set_formula(checks["K15"], '=IF(COUNT(Scenario_Analysis!$B$5:$E$9)<>20,1,0)', FMT_NOK)
    set_formula(
        checks["K16"],
        f'=IF(OR(ABS(MIN(Scenario_Analysis!$B$41:$F$41))>{tolerance},ABS(MAX(Scenario_Analysis!$B$41:$F$41))>{tolerance},ABS(MIN(Scenario_Analysis!$B$71:$F$71))>{tolerance},ABS(MAX(Scenario_Analysis!$B$71:$F$71))>{tolerance},ABS(MIN(Scenario_Analysis!$B$101:$F$101))>{tolerance},ABS(MAX(Scenario_Analysis!$B$101:$F$101))>{tolerance},ABS(MIN(Scenario_Analysis!$B$131:$F$131))>{tolerance},ABS(MAX(Scenario_Analysis!$B$131:$F$131))>{tolerance}),1,0)',
        FMT_NOK,
    )
    set_formula(
        checks["K17"],
        f'=IF(OR(ABS(INDEX(Scenario_Analysis!$B$5:$E$5,1,MATCH(Assumptions_and_Inputs!$B$4,Assumptions_and_Inputs!$B$5:$E$5,0))-Financial_Statements!$I$6)>{tolerance},ABS(INDEX(Scenario_Analysis!$B$6:$E$6,1,MATCH(Assumptions_and_Inputs!$B$4,Assumptions_and_Inputs!$B$5:$E$5,0))-Financial_Statements!$I$12)>{tolerance},ABS(INDEX(Scenario_Analysis!$B$8:$E$8,1,MATCH(Assumptions_and_Inputs!$B$4,Assumptions_and_Inputs!$B$5:$E$5,0))-Financial_Statements!$I$25)>{tolerance},ABS(INDEX(Scenario_Analysis!$B$9:$E$9,1,MATCH(Assumptions_and_Inputs!$B$4,Assumptions_and_Inputs!$B$5:$E$5,0))-Financial_Statements!$I$31)>{tolerance}),1,0)',
        FMT_NOK,
    )
    set_formula(checks["B2"], '=SUM(K6:K17)', FMT_NOK)
    checks["A2"] = '=IF(B2=0,"ALL CHECKS PASSED","MODEL CHECKS FAILED")'
    checks["A2"].fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    checks["A2"].font = Font(color="006100", bold=True)
    checks.conditional_formatting.add("K6:K17", ColorScaleRule(start_type="min", start_color="E2EFDA", end_type="max", end_color="FCE4D6"))
    checks.conditional_formatting.add("A2", FormulaRule(formula=["$B$2=0"], fill=PatternFill("solid", fgColor=LIGHT_GREEN), font=Font(color="006100", bold=True)))
    checks.conditional_formatting.add("A2", FormulaRule(formula=["$B$2>0"], fill=PatternFill("solid", fgColor=LIGHT_RED), font=Font(color="9C0006", bold=True)))

    # Executive summary and navigation.
    title(summary, "Amok Equipment AS | Integrated 3-Statement Model", 9)
    summary["A4"] = "Active scenario"
    set_formula(summary["B4"], "=Assumptions_and_Inputs!$B$4", "General")
    summary["D4"] = "Forecast period"
    summary["E4"] = "2026-2030"
    summary["G4"] = "Model status"
    set_formula(summary["H4"], "=Model_Checks!$A$2", "General")
    summary["A6"] = "Selected-scenario financial overview (NOK)"
    summary.merge_cells("A6:I6")
    summary["A6"].fill = PatternFill("solid", fgColor=BLUE)
    summary["A6"].font = Font(color=WHITE, bold=True)
    for col, year in enumerate(YEARS, start=2):
        summary.cell(7, col, year)
    summary["A7"] = "Metric"
    header_row(summary, 7, 1, 9)
    summary_metrics = [(8, "Revenue", 6), (9, "EBITDA", 12), (10, "EBITDA margin", 13), (11, "Net income", 20), (12, "Ending cash", 55), (13, "Total debt", 31)]
    for row, label, model_row in summary_metrics:
        summary.cell(row, 1, label)
        for year in YEARS:
            col = COLS[year]
            fmt = FMT_PCT if model_row == 13 and label.endswith("margin") else FMT_NOK
            set_formula(summary.cell(row, col), f"=Financial_Statements!{chr(col + 64)}{model_row}", fmt)
    summary["A16"] = "Scenario comparison | 2030 outcomes"
    summary.merge_cells("A16:G16")
    summary["A16"].fill = PatternFill("solid", fgColor=BLUE)
    summary["A16"].font = Font(color=WHITE, bold=True)
    summary["A17"] = "Scenario"
    summary["B17"] = "Revenue"
    summary["C17"] = "EBITDA"
    summary["D17"] = "EBITDA margin"
    summary["E17"] = "Ending cash"
    summary["F17"] = "Revolver debt"
    summary["G17"] = "Liquidity"
    header_row(summary, 17, 1, 7)
    for row, scenario_col, label in zip(range(18, 22), range(2, 6), SCENARIOS):
        summary.cell(row, 1, label)
        for out_col, analysis_row in [(2, 5), (3, 6), (4, 7), (5, 8), (6, 9), (7, 10)]:
            fmt = FMT_PCT if analysis_row == 7 else FMT_NOK
            set_formula(summary.cell(row, out_col), f"=Scenario_Analysis!{chr(scenario_col + 64)}{analysis_row}", fmt if analysis_row != 10 else None)
    summary["A24"] = "Workbook navigation"
    summary["A24"].fill = PatternFill("solid", fgColor=BLUE)
    summary["A24"].font = Font(color=WHITE, bold=True)
    nav = [
        ("Assumptions_and_Inputs", "Edit scenario selector and forecast assumptions"),
        ("Financial_Statements", "Integrated income statement, balance sheet, and cash flow"),
        ("Supporting_Schedules", "Working capital, PP&E, debt, and equity corkscrews"),
        ("Scenario_Analysis", "Compare scenario revenue and EBITDA"),
        ("Sensitivity_Analysis", "Two-way heatmap, tornado shocks, and elasticity"),
        ("Model_Checks", "Accounting and formula-integrity checks"),
        ("Historical_Data", "Source inputs and mapping notes"),
    ]
    for row, (sheet, description) in enumerate(nav, start=25):
        summary.cell(row, 1, sheet)
        summary.cell(row, 1).hyperlink = f"#'{sheet}'!A1"
        summary.cell(row, 1).style = "Hyperlink"
        summary.cell(row, 2, description)
    summary.merge_cells("A34:I34")
    summary["A34"] = "Blue fill / blue font = editable input; green font = formula. Base-case forecast assumptions are illustrative and require management review."
    summary["A34"].alignment = Alignment(wrap_text=True)
    summary.row_dimensions[34].height = 30

    chart = LineChart()
    chart.title = "Revenue and EBITDA (NOK)"
    chart.style = 13
    chart.y_axis.title = "NOK"
    chart.x_axis.title = "Fiscal year"
    chart.height = 8
    chart.width = 16
    chart.add_data(Reference(summary, min_col=1, max_col=9, min_row=8, max_row=9), titles_from_data=True, from_rows=True)
    chart.set_categories(Reference(summary, min_col=2, max_col=9, min_row=7))
    chart.legend.position = "b"
    chart.display_blanks = "gap"
    summary.add_chart(chart, "A37")

    # Make sheet widths readable and repeat headers for printing.
    widths = {
        "TOC_and_Summary": {"A": 43, "B": 20, "C": 18, "D": 20, "E": 20, "F": 16, "G": 17, "H": 19, "I": 15},
        "Assumptions_and_Inputs": {"A": 54, "B": 18, "C": 18, "D": 18, "E": 18, "F": 20},
        "Financial_Statements": {"A": 52, **{chr(col + 64): 15 for col in range(2, 10)}},
        "Supporting_Schedules": {"A": 58, **{chr(col + 64): 15 for col in range(2, 10)}},
        "Scenario_Analysis": {"A": 49, "B": 19, "C": 19, "D": 19, "E": 19, "F": 19},
        "Sensitivity_Analysis": {"A": 46, "B": 19, "C": 19, "D": 19, "E": 19, "F": 19, "G": 17, "H": 17},
        "Model_Checks": {"A": 36, "B": 68, **{chr(col + 64): 14 for col in range(3, 12)}},
        "Historical_Data": {"A": 58, "B": 18, "C": 18, "D": 18, "E": 18},
    }
    for ws in wb.worksheets:
        for col, width in widths[ws.title].items():
            ws.column_dimensions[col].width = width
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.print_title_rows = "1:5"
        ws.auto_filter.ref = None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and cell.row not in (1, 2, 3, 4, 5):
                    cell.alignment = Alignment(vertical="center", wrap_text=cell.alignment.wrap_text)
        if ws.title in {"Financial_Statements", "Supporting_Schedules", "Model_Checks"}:
            ws.freeze_panes = "B5"
        else:
            ws.freeze_panes = "B5"

    for ws in (statements, schedules, summary, scenario_analysis, sensitivity):
        ws.sheet_properties.outlinePr.summaryBelow = True
    statements.print_area = "A1:I59"
    schedules.print_area = "A1:I32"
    sensitivity.print_area = "A1:H23"
    checks.print_area = "A1:K17"
    scenario_analysis.print_area = "A1:F132"
    summary.print_area = "A1:I55"
    assumptions.print_area = "A1:F28"
    historical.print_area = "A1:E38"

    # Number formats for model schedule and statements.
    for ws in (statements, schedules):
        for row in range(5, ws.max_row + 1):
            for col in range(2, 10):
                cell = ws.cell(row, col)
                if cell.value is not None and cell.number_format == "General":
                    cell.number_format = FMT_NOK
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 85 if ws.title in {"Financial_Statements", "Supporting_Schedules", "Model_Checks", "Historical_Data"} else 90

    wb.save(OUTPUT_FILE)

    # Re-open and validate workbook structure/formula presence before reporting success.
    verified = load_workbook(OUTPUT_FILE, data_only=False)
    expected_sheets = [
        "TOC_and_Summary", "Assumptions_and_Inputs", "Financial_Statements",
        "Supporting_Schedules", "Scenario_Analysis", "Sensitivity_Analysis",
        "Model_Checks", "Historical_Data",
    ]
    assert verified.sheetnames == expected_sheets
    assert verified["Financial_Statements"]["I35"].data_type == "f"
    assert verified["Sensitivity_Analysis"]["E9"].data_type == "f"
    assert verified["Model_Checks"]["K14"].data_type == "f"
    assert verified["Assumptions_and_Inputs"]["B4"].value == "Base Case"
    assert OUTPUT_FILE.exists() and OUTPUT_FILE.stat().st_size > 10_000
    print(f"Created {OUTPUT_FILE.name} ({OUTPUT_FILE.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
